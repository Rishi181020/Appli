"""Word (.docx) resumes: the person's own file is the template, and every edit keeps its layout.

    skills tailoring (tailor.py)   one skill appended per suggestion to a Skills line, in that line's last run's font
    bank resume (bankbuild.py)     the entries of Experience / Projects / ... are replaced with the ones picked from the
                                   experience bank, each cloned from that section's own first entry (heading lines,
                                   bullet paragraph, spacing); the Skills lines are rewritten the same way. The header,
                                   Education and any section the bank has nothing for stay exactly as they are.
"""
import copy
import math
import re
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .tex import Unit

_SKILLS_HEADING = re.compile(r"\bskills?\b|core competencies|technologies|tech stack|tools", re.I)
_BULLET = re.compile(r"^\s*[•●▪‣◦\-–*]\s+")


def _paragraphs(doc) -> list[Paragraph]:
    """Every paragraph in reading order, table cells included (many resume templates lay out with tables)."""
    return [Paragraph(p, doc) for p in doc.element.body.iter(qn("w:p"))]


def _is_heading(p: Paragraph) -> bool:
    t = p.text.strip()
    if not t or len(t) > 45 or ":" in t or "," in t:
        return False
    runs = [r for r in p.runs if r.text.strip()]
    return (p.style is not None and p.style.name.lower().startswith(("heading", "title"))) or t.isupper() or \
        bool(runs and all(r.bold for r in runs))


def _is_bullet(p: Paragraph) -> bool:
    ppr = p._p.pPr  # bullets come from the paragraph's numbering, a list style ("List Bullet") or a typed character
    return (ppr is not None and ppr.numPr is not None) or bool(_BULLET.match(p.text)) or \
        (p.style is not None and p.style.name.lower().startswith("list"))


def _skill_line(text: str) -> tuple[str, str] | None:
    """'Languages: Python, Java' -> ('Languages', 'Python, Java'); a bare comma list -> ('Skills', list)."""
    label, sep, items = text.partition(":")
    if sep and items.strip() and len(label.strip()) <= 40:
        return _BULLET.sub("", label).strip(), items.strip()
    if text.count(",") >= 2:
        return "Skills", _BULLET.sub("", text).strip()
    return None


def parse(path: Path) -> dict:
    """-> {units, text, skills_text, bullets_text}: the Skills lines as editable units, plus the plain text."""
    doc = Document(str(path))
    paras = _paragraphs(doc)
    units, in_skills, bullets = [], False, []
    for i, p in enumerate(paras):
        text = p.text.strip()
        if _is_heading(p):
            in_skills = bool(_SKILLS_HEADING.search(text))
            continue
        if in_skills and text and any(r.text.strip() for r in p.runs):
            line = _skill_line(text)
            if line:
                units.append(Unit(f"w{i}", "skill", i, i, line[1], line[0]))
        elif _is_bullet(p) and text:
            bullets.append(_BULLET.sub("", text))
    return {
        "units": units,
        "text": "\n".join(p.text for p in paras),
        "skills_text": " ".join(u.text for u in units),
        "bullets_text": " ".join(bullets),
    }


def compose(items: str, skills: list[str]) -> str:
    """The Skills line with extra items appended, in its own comma-separated style."""
    return items.rstrip().rstrip(".") + "".join(f", {s}" for s in skills)


def apply(path: Path, additions: dict[str, list[str]]) -> tuple[bytes, str]:
    """Append skills to Skills lines (unit id -> skills) -> (the new .docx, its plain text)."""
    doc = Document(str(path))
    paras = _paragraphs(doc)
    for uid, skills in additions.items():
        p = paras[int(uid[1:])]
        run = next(r for r in reversed(p.runs) if r.text.strip())
        stripped = run.text.rstrip()
        trailing = run.text[len(stripped):]  # keep a tab / space the template ends the line with
        if stripped.endswith("."):
            stripped = stripped[:-1]
        run.text = stripped + "".join(f", {s}" for s in skills) + trailing
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), "\n".join(p.text for p in paras)


# ---- a whole resume from the experience bank, in the person's own layout ------------------------------------------
# section type -> what its heading says; first match wins ("Research Experience" is research) - same as bankbuild
SECTION_TYPES = [("education", r"educat"), ("publications", r"publicat|papers"), ("research", r"research"),
                 ("projects", r"project"), ("experience", r"experience|employment|\bwork\b"),
                 ("leadership", r"leadership|activit|involvement|volunteer"), ("awards", r"award|honou?r|achievement"),
                 ("skills", r"skill")]
_ANY_SECTION = re.compile(r"educat|experience|employment|\bwork\b|project|research|publicat|skill|summary|objective|"
                          r"profile|certific|interest|leadership|activit|volunteer|award|honou?r|course|languages", re.I)
_DATED = ("experience", "research", "leadership")


def _blocks(doc) -> list:
    """The body's top-level paragraphs and tables, in order."""
    out = []
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            out.append(Paragraph(el, doc))
        elif el.tag == qn("w:tbl"):
            out.append(Table(el, doc))
    return out


def _look(p: Paragraph) -> tuple:
    run = next((r for r in p.runs if r.text.strip()), None)
    return (p.style.name if p.style is not None else "", p.text.strip().isupper(),
            all(r.bold for r in p.runs if r.text.strip()), run.font.size if run is not None else None)


def _section_heads(blocks: list) -> set:
    """Section headings: short heading-like lines formatted like the ones that name a section ("Education",
    "EXPERIENCE"). A bold company name is formatted differently, so it isn't taken for a section."""
    named = [b for b in blocks if isinstance(b, Paragraph) and _is_heading(b) and _ANY_SECTION.search(b.text)]
    if not named:
        return set()
    looks = [_look(p) for p in named]
    common = max(set(looks), key=looks.count)
    return {id(b._p) for b in blocks if isinstance(b, Paragraph) and _is_heading(b) and _look(b) == common}


def _type(title: str) -> str | None:
    return next((t for t, pat in SECTION_TYPES if re.search(pat, title, re.I)), None)


def _sections(doc) -> list[dict]:
    """[{heading, type, body}] in order; the first one (heading None) is the header above every section."""
    blocks = _blocks(doc)
    heads = _section_heads(blocks)
    out = [{"heading": None, "type": None, "body": []}]
    for b in blocks:
        if isinstance(b, Paragraph) and id(b._p) in heads:
            out.append({"heading": b, "type": _type(b.text), "body": []})
        else:
            out[-1]["body"].append(b)
    return out


def section_titles(path: Path) -> list[str]:
    """The person's own section headings, in order (for the planner: their wording and their order)."""
    return [s["heading"].text.strip() for s in _sections(Document(str(path)))[1:]]


def _entries(body: list) -> list[dict]:
    """A section's entries: heading lines, then bullets, then any spacer lines."""
    out, cur = [], None
    for b in body:
        is_p = isinstance(b, Paragraph)
        if is_p and not b.text.strip() and not b._p.xpath(".//w:drawing"):
            if cur is not None:
                cur["tail"].append(b)
            continue
        if is_p and _is_bullet(b):
            if cur is None:
                cur = {"head": [], "bullets": [], "tail": []}
                out.append(cur)
            cur["bullets"].append(b)
        else:
            if cur is None or cur["bullets"] or cur["tail"]:
                cur = {"head": [], "bullets": [], "tail": []}
                out.append(cur)
            cur["head"].append(b)
    return out


def _texts(el) -> list:
    return list(el.iter(qn("w:t")))


def _el_text(el) -> str:
    return "".join(t.text or "" for t in _texts(el))


def _dates(e: dict) -> str:
    return " - ".join(x for x in (e.get("start"), e.get("end")) if x)


def _source_entry(head_text: str, bank_entries: list[dict]) -> dict | None:
    """The bank entry the template's own first entry came from (its org / title / dates are in the heading lines)."""
    low = head_text.lower()
    best, score = None, 0
    for e in bank_entries:
        s = sum(len(v) for v in (e.get("org"), e.get("title")) if v and v.lower() in low)
        if s > score:
            best, score = e, s
    return best


def _sep_clean(s: str) -> str:
    """Separators left dangling after a field was emptied: 'Acme | ' -> 'Acme'."""
    s = re.sub(r"(\s*[|·•,–—-]\s*){2,}", lambda m: m.group(1), s)
    return re.sub(r"^\s*[|·•,–—-]\s*|\s*[|·•,–—-]\s*$", "", s) if s.strip(" |·•,–—-") else ""


def _fill_head(els: list, src: dict | None, e: dict, doc) -> None:
    """Rewrite cloned heading lines for entry `e`: each field of the template's own entry (org, title, location,
    start, end) is swapped for e's, wherever it is, so the formatting of every piece stays."""
    pairs = []
    if src:
        for f in ("org", "title", "location", "start", "end"):
            old = (src.get(f) or "").strip()
            if old and any(old.lower() in _el_text(el).lower() for el in els):
                pairs.append((old, (e.get(f) or "").strip()))
    if pairs:
        pat = re.compile("|".join(re.escape(o) for o, _ in sorted(pairs, key=lambda x: -len(x[0]))), re.I)
        new = {o.lower(): n for o, n in pairs}
        for el in els:
            for p in [el] if el.tag == qn("w:p") else el.iter(qn("w:p")):
                nodes = _texts(p)
                whole = "".join(t.text or "" for t in nodes)
                if not pat.search(whole):
                    continue
                if any(pat.search(t.text or "") for t in nodes):  # each field sits in its own run(s)
                    for t in nodes:
                        if t.text and pat.search(t.text):
                            t.text = _sep_clean(pat.sub(lambda m: new[m.group(0).lower()], t.text)) or ""
                            t.set(qn("xml:space"), "preserve")
                else:  # a field split across runs: the line takes the first run's formatting
                    nodes[0].text = _sep_clean(pat.sub(lambda m: new[m.group(0).lower()], whole))
                    for t in nodes[1:]:
                        t.text = ""
    else:  # no known entry to map from: by position (line 1 = org / title + dates, line 2 = title + location)
        lines = [p for el in els for p in ([el] if el.tag == qn("w:p") else el.iter(qn("w:p"))) if _el_text(p).strip()]
        vals = [[e.get("org") or e.get("title") or "", _dates(e)], [e.get("title") if e.get("org") else "", e.get("location") or ""]]
        for p, (left, right) in zip(lines, vals):
            nodes = [t for t in _texts(p) if (t.text or "").strip()]
            if nodes:
                nodes[0].text = left
                for t in nodes[1:-1]:
                    t.text = ""
                if len(nodes) > 1:
                    nodes[-1].text = right
    # links: the template's link now points at this entry's own link; without one, a linked name stays as plain text
    # and a bare "[Link]" / "GitHub" goes
    fields = [v.lower() for v in (e.get(f) for f in ("org", "title", "location")) if v]
    for el in els:
        for h in list(el.iter(qn("w:hyperlink"))):
            if not h.get(qn("r:id")):
                continue
            if (e.get("link") or "").startswith("http"):
                h.set(qn("r:id"), doc.part.relate_to(e["link"], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            elif any(v in _el_text(h).lower() for v in fields):
                for child in list(h):
                    h.addprevious(child)
                h.getparent().remove(h)
            else:
                h.getparent().remove(h)


def _set_runs(p_el, text: str) -> None:
    """A cloned paragraph says `text`, in its first run's formatting; **bold** spans become bold runs."""
    p = Paragraph(p_el, None)
    runs = [r for r in p.runs if r.text.strip()] or p.runs
    if not runs:
        p.add_run(text)
        return
    proto = runs[0]._r
    for r in p.runs:
        r._r.getparent().remove(r._r)
    for i, piece in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not piece:
            continue
        r = copy.deepcopy(proto)
        for child in list(r):  # keep only the formatting: no old text, tabs or breaks
            if child.tag != qn("w:rPr"):
                r.remove(child)
        p_el.append(r)
        run = Paragraph(p_el, None).runs[-1]
        run.text = piece
        if i % 2:
            run.bold = True


def _skill_paragraph(proto_el, group: str, items: str):
    """A Skills line in the template line's own formatting: the label runs say `group:`, the item runs `items`."""
    el = copy.deepcopy(proto_el)
    p = Paragraph(el, None)
    runs = [r for r in p.runs if r.text]
    k = next((i for i, r in enumerate(runs) if ":" in r.text), None)
    if k is None:  # a bare list
        _set_runs(el, f"{group}: {items}" if group and group != "Skills" else items)
        return el
    lead = re.match(r"^\s*[•●▪‣◦\-–*]?\s*", runs[0].text).group(0)
    _, _, after = runs[k].text.partition(":")
    space = after[: len(after) - len(after.lstrip())] or " "
    items_here = bool(after.strip()) or k == len(runs) - 1  # the colon's run also holds the items
    colon = f":{space}{items}" if items_here else f":{after}"
    if k == 0:
        runs[0].text = f"{lead}{group}{colon}"
    else:  # the label's own run keeps its formatting (often bold); the colon stays in its run
        runs[0].text = f"{lead}{group}"
        for r in runs[1:k]:
            r.text = ""
        runs[k].text = colon
    rest = runs[k + 1:]
    if not items_here:
        rest[0].text = items
        rest = rest[1:]
    for r in rest:
        r.text = ""
    return el


def _cpl(doc, sample: Paragraph | None) -> float:
    """About how many characters fit on one line of body text."""
    s = doc.sections[0]
    width = (s.page_width - s.left_margin - s.right_margin) / 12700 if s.page_width else 540  # EMU -> pt
    size = None
    if sample is not None:
        run = next((r for r in sample.runs if r.text.strip()), None)
        size = (run.font.size if run is not None else None) or (sample.style.font.size if sample.style is not None else None)
    pt = size.pt if size else 11
    return max(40.0, width / (pt * 0.5))


def _lines(blocks: list, cpl: float) -> int:
    """Estimated printed lines: a paragraph wraps every `cpl` characters; a table row takes its tallest cell."""
    n = 0
    for b in blocks:
        if isinstance(b, Paragraph):
            n += max(1, math.ceil(len(b.text) / cpl))
        else:
            for row in b.rows:
                n += max([sum(max(1, math.ceil(len(p.text) / (cpl / max(1, len(row.cells))))) for p in c.paragraphs)
                          for c in row.cells] or [1])
    return n


def _pages_saved(path: Path) -> int | None:
    """The page count Word saved in the file (docProps/app.xml); None if it isn't there (e.g. a Google Docs export)."""
    try:
        with zipfile.ZipFile(path) as z:
            m = re.search(rb"<Pages>(\d+)</Pages>", z.read("docProps/app.xml"))
        return max(1, int(m.group(1))) if m else None
    except (KeyError, zipfile.BadZipFile, OSError):
        return None


def _lines_per_page(doc, sample: Paragraph | None) -> float:
    """Printed lines a page holds, from its height, margins and the body font (single spacing plus a little)."""
    s = doc.sections[0]
    height = (s.page_height - s.top_margin - s.bottom_margin) / 12700 if s.page_height else 720
    run = next((r for r in sample.runs if r.text.strip()), None) if sample is not None else None
    pt = run.font.size.pt if run is not None and run.font.size else 11
    return height / (pt * 1.3)


def bank_maker(path: Path, bank_entries: list[dict]):
    """-> make(plan, draft, by_id) -> (docx bytes, layout): the Word counterpart of render + compile for bankbuild's
    fit / fill. Word lays the page out itself, so the page count is an estimate: the person's own resume is the
    measure (its lines per page), so a 1-page resume stays about as long as theirs."""
    base = Document(str(path))
    secs = _sections(base)
    if not any(s["type"] in ("experience", "projects", "research") and _entries(s["body"]) for s in secs):
        raise ValueError("couldn't find the Experience / Projects entries in your Word resume")
    sample = next((b for s in secs for b in s["body"] if isinstance(b, Paragraph) and _is_bullet(b)), None)
    cpl = _cpl(base, sample)
    own = _lines(_blocks(base), cpl)
    pages = _pages_saved(path) or max(1, round(own / _lines_per_page(base, sample)))
    per_page = own / pages  # their resume, as they laid it out, is what one page holds

    def make(plan_: dict, draft: dict, by_id: dict):
        doc = Document(str(path))
        secs = _sections(doc)
        protos = {}  # section type -> (its first full entry, the bank entry it came from)
        for s in secs:
            ents = [x for x in _entries(s["body"]) if x["head"] and x["bullets"]]
            if s["type"] and ents:
                protos.setdefault(s["type"], (ents[0], _source_entry(" ".join(_el_text(b._element) for b in ents[0]["head"]), bank_entries)))
        heading_proto = next((s["heading"] for s in secs if s["heading"] is not None), None)

        def proto_for(t: str):
            if t in protos:
                return protos[t]
            similar = ("experience", "research", "leadership") if t in _DATED else ("projects", "awards", "publications")
            return next((protos[x] for x in (*similar, *protos) if x in protos), None)

        def entry_blocks(t: str, ids: list[dict]) -> list:
            proto = proto_for(t)
            if proto is None:
                return []
            ent, src = proto
            out = []
            for x in ids:
                e = by_id.get(x["id"])
                bullets = (draft.get("entries", {}).get(x["id"]) or {}).get("bullets", [])
                if not e or not bullets:
                    continue
                head = [copy.deepcopy(b._element) for b in ent["head"]]
                _fill_head(head, src, e, doc)
                out += head
                for b in bullets:
                    el = copy.deepcopy(ent["bullets"][0]._element)
                    _set_runs(el, b["text"])
                    out.append(el)
                out += [copy.deepcopy(b._element) for b in ent["tail"]]
            return out

        by_type = {s["type"]: s for s in secs if s["type"]}
        for ps in plan_["sections"]:
            t = ps["type"]
            if t == "education":
                continue  # theirs, as it is
            s = by_type.get(t)
            if t == "skills":
                lines = draft.get("skills") or plan_.get("base_skills") or []
                if s is None or not lines:
                    continue
                skill_ps = [b for b in s["body"] if isinstance(b, Paragraph) and b.text.strip() and _skill_line(b.text.strip())]
                if not skill_ps:
                    continue
                anchor = skill_ps[0]._element
                for g in lines:
                    anchor.addprevious(_skill_paragraph(skill_ps[0]._element, g["group"], g["items"]))
                for b in skill_ps:
                    b._element.getparent().remove(b._element)
                continue
            new = entry_blocks(t, ps["entries"])
            if s is not None:
                old = [b._element for b in s["body"]]
                anchor = s["heading"]._element
                for el in new:
                    anchor.addnext(el)
                    anchor = el
                if new:  # nothing picked for it: their own section stays as it was
                    for el in old:
                        el.getparent().remove(el)
            elif new and heading_proto is not None:  # a section their resume doesn't have (e.g. Publications)
                h = copy.deepcopy(heading_proto._element)
                _set_runs(h, ps["heading"].upper() if heading_proto.text.strip().isupper() else ps["heading"])
                at = by_type.get("skills")
                anchor_el = at["heading"]._element if at else doc.element.body[-1]
                for el in [h, *new]:
                    anchor_el.addprevious(el)
        buf = BytesIO()
        doc.save(buf)
        out = Document(BytesIO(buf.getvalue()))
        lines = _lines(_blocks(out), cpl)
        pages = max(1, math.ceil(lines / per_page - 0.02))
        layout = SimpleNamespace(ok=True, pages=pages, pdf=None, log="",
                                 fill=(lines - (pages - 1) * per_page) / per_page,
                                 text="\n".join(p.text for p in _paragraphs(out)))
        return buf.getvalue(), layout

    return make
