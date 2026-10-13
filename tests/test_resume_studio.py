"""Offline matching-to-download coverage; no model calls or application rows are created."""
import io
import tempfile
import unittest
import math
import pymupdf
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter
from runner import config
from runner.resume import manual, match, pdf, select


class ResumeStudioTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.home = self.root / "person-a"
        self.home.mkdir()
        self.source = self.home / "resume.docx"
        document = Document()
        document.add_paragraph("Example Candidate")
        document.add_heading("Skills", 1)
        document.add_paragraph("Languages: Python, Java, SQL")
        document.add_heading("Experience", 1)
        document.add_paragraph("Built reliable services.", style="List Bullet")
        document.save(self.source)
        self.addCleanup(select.base_docs.cache_clear)
        values = {"ROOT": self.root, "HOME": self.home, "PROFILE_PATH": self.home / "profile.md",
                  "PRIMARY_TEX": None, "RESUMES": {"base": self.source}, "RESUME_DOCX": {"base": self.source},
                  "RESUME_TEX": {}, "RESUME_LABELS": {"base": "My resume"}}
        self.enterContext(patch.multiple(config, **values))
        self.enterContext(patch("runner.resume.manual.profile_md.sync_from_db"))
        self.enterContext(patch("runner.resume.manual.load_profile", return_value=SimpleNamespace(
            text="I work with Docker and Python.", standard_answers={})))
        self.enterContext(patch("runner.resume.manual.active_screens", return_value=frozenset({"no_sponsor"})))
        self.enterContext(patch("runner.resume.build.file_stem", return_value="Example Resume Engineer"))
        self.enterContext(patch("runner.resume.tailor._place", side_effect=lambda skills, lines, jid: {t["term"]: lines[0].id for t in skills}))
        terms = [{"term": term, "aliases": [], "required": True, "category": "tool"}
                 for term in ("Python", "Java", "SQL", "Docker", "Kubernetes")]
        before = match.score("Python Java SQL", terms)
        self.info = {"chosen": "base", "pct": before["pct"], "terms": terms,
                     "missing": ["Docker", "Kubernetes"], "equivalent": [], "equivalent_any": [],
                     "tailorable": True, "format": "docx"}
        self.enterContext(patch("runner.resume.manual.select.evaluate", return_value=self.info))
        self.database = self.enterContext(patch("runner.resume.build.sb", side_effect=AssertionError("must not update application rows")))

    def analyze(self):
        return manual.analyze({"role": "Engineer", "description": "No visa sponsorship. " + "Python services and required engineering skills. " * 10})

    def test_review_export_and_private_download_file(self):
        draft = self.analyze()
        self.assertIn("No visa sponsorship", draft["restriction"])
        defaults = [e["id"] for e in draft["edits"]["edits"] if e["default"]]
        self.assertEqual(len(defaults), 1)
        result = manual.export(draft["draft_id"], defaults)
        self.assertTrue(result["ok"], result)
        self.assertGreater(result["pct_after"], draft["match"]["pct"])
        output = self.root / result["file"]
        self.assertTrue(output.is_relative_to(self.home))
        text = "\n".join(p.text for p in Document(io.BytesIO(output.read_bytes())).paragraphs)
        self.assertIn("Docker", text)
        self.assertNotIn("Kubernetes", text)
        self.assertIn("Built reliable services.", text)
        self.assertNotIn("Docker", "\n".join(p.text for p in Document(self.source).paragraphs))
        self.database.assert_not_called()
        # Another person's active folder cannot locate this draft.
        with patch.object(config, "HOME", self.root / "person-b"):
            with self.assertRaisesRegex(ValueError, "Draft not found"):
                manual.export(draft["draft_id"], defaults)

    def test_invalid_requests_stale_sources_and_unchanged_export(self):
        with self.assertRaises(ValueError):
            manual.analyze({"description": "too short"})
        with self.assertRaises(ValueError):
            manual.folder("../../person-a")
        draft = self.analyze()
        with self.assertRaisesRegex(ValueError, "Unknown suggested change"):
            manual.export(draft["draft_id"], ["invented"])
        result = manual.export(draft["draft_id"], [])
        self.assertTrue(result["ok"])
        self.assertTrue(result["unchanged"])
        document = Document(self.source)
        document.add_paragraph("New experience")
        document.save(self.source)
        with self.assertRaisesRegex(ValueError, "base resumes changed"):
            manual.export(draft["draft_id"], [])

    def test_pdf_only_review_build_and_content_preservation(self):
        source = self.home / "original.pdf"
        source_text = "Example Candidate\nexample@example.com | 555-0100\nSkills\nPython, Java, SQL\nExperience\nBuilt reliable services and reduced costs by 25%.\nEducation\nExample University, 2024"
        source.write_bytes(pdf.render(source_text, []))
        original = source.read_bytes()
        self.info.update(tailorable=False, format="tex")
        with patch.multiple(config, RESUMES={"base": source}, RESUME_DOCX={}):
            draft = self.analyze()
            self.assertEqual(draft["edits"]["format"], "pdf_rebuild")
            self.assertNotIn("no_latex", str(draft["edits"]["problems"]))
            defaults = [e["id"] for e in draft["edits"]["edits"] if e["default"]]
            with self.assertRaisesRegex(ValueError, "Review and confirm"):
                manual.export(draft["draft_id"], defaults)
            result = manual.export(draft["draft_id"], defaults, extraction_confirmed=True)
            self.assertTrue(result["ok"], result)
            self.assertTrue(result["preserved_pdf_layout"])
            self.assertEqual(result["pages"], 1)
            self.assertGreater(result["pct_after"], self.info["pct"])
            reader = PdfReader(self.root / result["file"])
            text = "\n".join(page.extract_text() for page in reader.pages)
            self.assertIn("Docker", text)
            self.assertNotIn("Kubernetes", text)
            self.assertIn("reduced costs by 25%", text)
            self.assertIn("Example University, 2024", text)
            self.assertEqual(source.read_bytes(), original)
            with pymupdf.open(source) as before, pymupdf.open(self.root / result["file"]) as after:
                region = pdf.skills_layout(source)["rect"]
                first, second = before[0].get_pixmap(), after[0].get_pixmap()
                stride = first.stride
                self.assertEqual(first.samples[:math.floor(region[1]) * stride], second.samples[:math.floor(region[1]) * stride])
                self.assertEqual(first.samples[math.ceil(region[3]) * stride:], second.samples[math.ceil(region[3]) * stride:])
            self.database.assert_not_called()

    def test_pdf_overflow_rejected_and_original_export_is_identical(self):
        source = self.home / "tight.pdf"
        source.write_bytes(pdf.render("Example Candidate\nSkills\nPython, SQL, Java\nExperience\nBuilt reliable services and reduced costs.", []))
        self.assertEqual(pdf.preserve_layout(source, []), source.read_bytes())
        with self.assertRaisesRegex(ValueError, "do not fit"):
            pdf.preserve_layout(source, [{"skill": "Very long skill phrase " * 80, "category": "tool"}])
        self.assertEqual(len(PdfReader(source).pages), 1)
        multi = self.home / "multi.pdf"
        multi.write_bytes(pdf.render("Example Candidate\nSkills\nPython, SQL, Java\nExperience\n" + "Long experience detail.\n" * 90, []))
        with self.assertRaisesRegex(ValueError, "more than one page"):
            pdf.skills_layout(multi)

    def test_serif_font_skills_have_readable_unicode(self):
        source = self.home / "serif.pdf"
        with pymupdf.open() as document:
            page = document.new_page(width=612, height=792)
            page.insert_text((40, 40), "Example Candidate", fontname="tiro", fontsize=16)
            page.insert_text((40, 80), "Skills", fontname="tibo", fontsize=12)
            page.insert_text((40, 100), "Tools: Python, SQL, Git", fontname="tiro", fontsize=10)
            page.insert_text((40, 165), "Experience", fontname="tibo", fontsize=12)
            page.insert_text((40, 185), "Built reliable services.", fontname="tiro", fontsize=10)
            document.save(source)
        data = pdf.preserve_layout(source, [{"skill": skill, "category": "tool"}
                                           for skill in ("Kafka", "AWS", "Docker", "REST API design")])
        with self.assertNoLogs("pypdf", level="WARNING"):
            text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
        for skill in ("Kafka", "AWS", "Docker", "REST API design"):
            self.assertIn(skill, text.replace("\u00a0", " "))

    def test_unreadable_scanned_page_and_escaped_long_text(self):
        source = self.home / "scanned.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.write(source)
        with self.assertRaisesRegex(ValueError, "scanned or unreadable"):
            pdf.extract(source)
        text = "Example Candidate\nSkills\nC++, C#, R&D\nExperience\n" + "Used <tools> & services to cut costs by 20%.\n" * 90
        data = pdf.render(text, ["Docker"])
        reader = PdfReader(io.BytesIO(data))
        extracted = "\n".join(page.extract_text() for page in reader.pages)
        self.assertGreater(len(reader.pages), 1)
        self.assertIn("C++, C#, R&D", extracted)
        self.assertEqual(extracted.count("Used <tools> & services to cut costs by 20%."), 90)


if __name__ == "__main__":
    unittest.main()
