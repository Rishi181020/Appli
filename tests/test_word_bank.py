"""A resume built from the experience bank for a Word resume keeps the person's own layout: header, Education and
headings untouched; Experience / Projects entries swapped for the bank's picks in the template entry's formatting;
Skills lines rewritten in theirs; the page-length estimate trims to fit."""
import tempfile
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from runner.resume import bankbuild, word

d = Path(tempfile.mkdtemp())
doc = Document()
doc.add_paragraph().add_run("Jane Doe").bold = True
doc.add_paragraph("jane@example.com | (555) 010-0000")


def heading(text):
    doc.add_paragraph().add_run(text).bold = True


def entry_head(org, dates, title, loc):
    p = doc.add_paragraph()
    p.add_run(org).bold = True
    p.add_run("\t")
    p.add_run(dates)
    p = doc.add_paragraph()
    p.add_run(title).italic = True
    p.add_run("\t")
    p.add_run(loc).italic = True


heading("EDUCATION")
doc.add_paragraph("University of Illinois Urbana-Champaign, B.S. Statistics, May 2025")
heading("EXPERIENCE")
entry_head("Acme Analytics", "Jun 2024 - Present", "Data Analyst Intern", "Chicago, IL")
doc.add_paragraph("Built weekly KPI dashboards in Tableau", style="List Bullet")
doc.add_paragraph("Cleaned 2M rows of sales data with SQL", style="List Bullet")
entry_head("Campus Library", "Jan 2023 - May 2024", "Student Assistant", "Urbana, IL")
doc.add_paragraph("Helped 40 students a day", style="List Bullet")
for i in range(14):  # a real resume fills its page: this one is about as long as one
    doc.add_paragraph(f"Shelved and catalogued returns for the reference desk, batch {i}, and answered questions", style="List Bullet")
heading("PROJECTS")
p = doc.add_paragraph()
p.add_run("Course Planner").bold = True
link = OxmlElement("w:hyperlink")
link.set(qn("r:id"), doc.part.relate_to("https://github.com/jane/planner", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
r = OxmlElement("w:r")
t = OxmlElement("w:t")
t.text = " [Link]"
r.append(t)
link.append(r)
p._p.append(link)
doc.add_paragraph("Planned 4-year schedules with Python", style="List Bullet")
heading("TECHNICAL SKILLS")
for label, items in (("Languages", "Python, SQL, R"), ("Tools", "Tableau, Excel")):
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(items)
for para in doc.paragraphs:
    for run in para.runs:
        run.font.size = Pt(11)
doc.save(d / "jane.docx")

bank = [
    {"id": "acme", "kind": "experience", "org": "Acme Analytics", "title": "Data Analyst Intern", "location": "Chicago, IL",
     "start": "Jun 2024", "end": "Present", "link": ""},
    {"id": "lib", "kind": "experience", "org": "Campus Library", "title": "Student Assistant", "location": "Urbana, IL",
     "start": "Jan 2023", "end": "May 2024", "link": ""},
    {"id": "bank1", "kind": "experience", "org": "Northwind Health", "title": "Business Intelligence Intern",
     "location": "Remote", "start": "May 2023", "end": "Aug 2023", "link": ""},
    {"id": "plan", "kind": "project", "org": "", "title": "Course Planner", "location": "", "start": "", "end": "",
     "link": "https://github.com/jane/planner"},
    {"id": "proj2", "kind": "project", "org": "", "title": "Churn Model", "location": "", "start": "", "end": "", "link": ""},
]
by_id = {e["id"]: e for e in bank}
B = lambda *xs: [{"text": x, "evidence": "", "terms": []} for x in xs]  # noqa: E731
draft = {"entries": {"acme": {"bullets": B("Automated **Power BI** reports for 6 teams", "Cut query time 40% with SQL indexes")},
                     "bank1": {"bullets": B("Built Azure Data Factory pipelines for claims data")},
                     "proj2": {"bullets": B("Predicted churn with logistic regression in Python")}},
         "skills": [{"group": "Languages", "items": "SQL, Python"}, {"group": "Tools", "items": "Power BI, Azure Data Factory, Tableau"}]}
plan = {"pages": 1, "sections": [
    {"type": "education", "heading": "EDUCATION", "entries": []},
    {"type": "experience", "heading": "EXPERIENCE", "entries": [{"id": "acme", "bullets": 2}, {"id": "bank1", "bullets": 1}]},
    {"type": "projects", "heading": "PROJECTS", "entries": [{"id": "proj2", "bullets": 1}]},
    {"type": "skills", "heading": "TECHNICAL SKILLS", "entries": []}], "reserve": []}

assert word.section_titles(d / "jane.docx") == ["EDUCATION", "EXPERIENCE", "PROJECTS", "TECHNICAL SKILLS"]
data, layout = word.bank_maker(d / "jane.docx", bank)(plan, draft, by_id)
(d / "out.docx").write_bytes(data)
out = Document(str(d / "out.docx"))
texts = [p.text for p in out.paragraphs]

# header and Education exactly as they were
assert texts[:4] == ["Jane Doe", "jane@example.com | (555) 010-0000", "EDUCATION",
                     "University of Illinois Urbana-Champaign, B.S. Statistics, May 2025"], texts[:4]
# Experience: the bank's picks, newest first as planned, in the template entry's lines
assert "Acme Analytics\tJun 2024 - Present" in texts and "Northwind Health\tMay 2023 - Aug 2023" in texts, texts
assert "Business Intelligence Intern\tRemote" in texts
assert "Campus Library\tJan 2023 - May 2024" not in texts, "an entry the plan left out is gone"
nw = next(p for p in out.paragraphs if p.text.startswith("Northwind"))
assert nw.runs[0].bold and not nw.runs[-1].bold, "company bold, dates plain: the template's formatting"
bullet = next(p for p in out.paragraphs if p.text.startswith("Automated"))
assert bullet.style.name == "List Bullet" and bullet.text == "Automated Power BI reports for 6 teams"
assert [r.bold for r in bullet.runs if r.text == "Power BI"] == [True], "**bold** in a bullet becomes a bold run"
# Projects: the new project, and the template's [Link] goes since it has no link
assert "Churn Model" in texts and "Course Planner" not in texts
assert not next(p for p in out.paragraphs if p.text.startswith("Churn"))._p.xpath(".//w:hyperlink")
# Skills: their labels in bold, the new items plain
sk = next(p for p in out.paragraphs if p.text.startswith("Tools"))
assert sk.text == "Tools: Power BI, Azure Data Factory, Tableau" and sk.runs[0].bold and not sk.runs[-1].bold, sk.text
assert layout.pages == 1 and 0 < layout.fill <= 1.05 and "Northwind" in layout.text

# fit: with far too many bullets for their one page, the least relevant ones are trimmed until it fits
many = {k: {"bullets": B(*[f"Delivered measurable outcome number {i} for the analytics team using SQL" for i in range(12)])}
        for k in ("acme", "bank1", "proj2")}
big = {"entries": many, "skills": draft["skills"]}
src, comp, trimmed = bankbuild.fit(plan, big, by_id, word.bank_maker(d / "jane.docx", bank))
assert comp.pages == 1 and trimmed, (comp.pages, len(trimmed))
print("all word bank checks passed")
