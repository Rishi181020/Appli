"""Word resumes: the Skills lines are found (body and table layouts), skills are appended in the line's own font, and
nothing else in the file changes."""
import tempfile
from pathlib import Path

from docx import Document

from runner.resume import word

d = Path(tempfile.mkdtemp())

# a common layout: bold section headings, "Label: items" lines with a bold label run, bullets under experience
doc = Document()
doc.add_paragraph().add_run("EXPERIENCE").bold = True
doc.add_paragraph("Built ETL pipelines in Python and SQL", style="List Bullet")
doc.add_paragraph().add_run("TECHNICAL SKILLS").bold = True
p = doc.add_paragraph()
p.add_run("Languages: ").bold = True
p.add_run("Python, SQL, R.")
p = doc.add_paragraph()
p.add_run("Tools: ").bold = True
p.add_run("Tableau, Excel")
doc.add_paragraph().add_run("EDUCATION").bold = True
doc.add_paragraph("University of Illinois Urbana-Champaign, B.S. Statistics")
doc.save(d / "body.docx")

w = word.parse(d / "body.docx")
assert [(u.context, u.text) for u in w["units"]] == [("Languages", "Python, SQL, R."), ("Tools", "Tableau, Excel")], w["units"]
assert "ETL pipelines" in w["bullets_text"] and "Tableau" in w["skills_text"]
langs, tools = w["units"]
data, text = word.apply(d / "body.docx", {langs.id: ["Power BI"], tools.id: ["Azure Data Factory", "Power Automate"]})
assert "Languages: Python, SQL, R, Power BI" in text, text
assert "Tools: Tableau, Excel, Azure Data Factory, Power Automate" in text, text
assert word.compose(langs.text, ["Power BI"]) == "Python, SQL, R, Power BI"  # what Resume review shows = what's written
(d / "out.docx").write_bytes(data)
out = Document(str(d / "out.docx"))
line = next(p for p in out.paragraphs if p.text.startswith("Languages"))
assert line.runs[0].bold and not line.runs[-1].bold, "added skills take the items' font, the label stays bold"
assert [p.text for p in out.paragraphs if not p.text.startswith(("Languages", "Tools"))] == \
       [p.text for p in Document(str(d / "body.docx")).paragraphs if not p.text.startswith(("Languages", "Tools"))]

# a table layout (two-column templates) with a styled heading
doc = Document()
doc.add_heading("Skills", level=1)
t = doc.add_table(rows=1, cols=1)
t.cell(0, 0).paragraphs[0].add_run("Cloud: AWS, GCP, Docker")
doc.save(d / "table.docx")
w = word.parse(d / "table.docx")
assert [u.context for u in w["units"]] == ["Cloud"], w["units"]
_, text = word.apply(d / "table.docx", {w["units"][0].id: ["Kubernetes"]})
assert "Cloud: AWS, GCP, Docker, Kubernetes" in text

print("all word resume checks passed")
